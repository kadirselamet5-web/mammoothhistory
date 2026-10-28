#!/usr/bin/env python3
"""Build the Mammoth History Lab discovery map (single-file HTML) from the Lab's workbook.

Usage:
    python3 scripts/build_atlas.py                      # uses the defaults below
    python3 scripts/build_atlas.py --workbook data/NEW.xlsx --out index.html --previous index.html

The script never rewrites a workbook cell. Every original cell value is kept in each record's
``source_values``; the display fields next to them are derived (and documented) here so the
interface can filter, sort and plot. Rules of thumb used throughout:

* A cell is a *value* unless it is blank, a "not applicable / not available / unknown" marker
  (status ``na``) or a bare question mark (status ``q``).
* Years are read from the first four-digit number in a date cell ("1882-1884" -> 1882).
* Coordinates must be numeric; text cells are parsed when they contain a plain number and the
  record receives a review note. Records without usable coordinates stay in the list and table.
* Material tags (facets) are keyword browsing aids derived from ``Parts_found``; negated phrases
  ("missing trunk and tail", "no head") are ignored before tagging.
* Completeness grade follows the ``Complete?`` column.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html
import json
import re
import sys
import warnings
from collections import Counter, OrderedDict
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT / "data" / "SOFT_TISSUE_Project_Database_2026_10_28_AO.xlsx"
DEFAULT_OUT = ROOT / "index.html"
TEMPLATE = ROOT / "src" / "atlas.html"
CSS = ROOT / "src" / "atlas.css"
APP = ROOT / "src" / "atlas.js"
LAND = ROOT / "src" / "land.txt"
THREE = ROOT / "vendor" / "three.r128.min.js"
REVIEW_NOTES = ROOT / "data" / "review_notes.json"

# Columns, by header, that the interface reads. Older workbooks used slightly different headers.
HEADER_ALIASES = {
    "Secondary_name": "Secondary_name(s)",
    "Location_Provenance": "Location Provenance",
    "Coordinate_provenance": "Location Provenance",
}
NA_MARKERS = {"not_applicable", "not_available", "n/a", "na", "none", "unknown", "undated", "not recorded", "not_recorded", "-"}
GRADES = {"Fragments_Only": 1, "Partial": 2, "Mostly_Complete": 3, "Complete": 4}

# Material tags. Order matters only for display. Patterns run on a lower-cased description from
# which negated phrases have been removed. Keep the vocabulary in sync with MATERIALS in atlas.js.
FACET_RULES = OrderedDict([
    ("skin", r"\bskins?\b|\bhide\b"),
    ("hair", r"\bhairs?\b|\bfur\b|\bwool\b|\bhaior\b|\bwhiskers?\b|\beyelash"),
    ("flesh", r"\bflesh\b|\bmuscles?\b|\bmuslce\b|\btongue\b"),
    ("organs", r"\borgans?\b|\bbrains?\b|\bintestin|\bstomach\b|\bgastro|\btrachea\b|\bdiaphra|\bliver\b|\bheart\b|\blungs?\b|\bgland\b"),
    ("trunk", r"\btrunks?\b"),
    ("connective", r"\bligaments?\b|\btendons?\b"),
    ("keratin", r"\bhoofs?\b|\bhooves\b|\bclaws?\b|\bnails?\b|\bhorns?\b|\bsheaths?\b"),
    ("softgeneric", r"soft[\s-]*(?:tissues?|itssues?|parts)|\bmummified\b|\bmummy\b"),
])
NEGATION = re.compile(r"(?:\*+\s*)?\b(?:missing|no|without|lacking|absent|except)\b[^,;.]*", re.I)
PROVENANCE_ORDER = ["source-stated", "historical map", "inferred/historical map", "inferred", "regional_placeholder", "unlocated"]


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(text).strip().lower()).strip("_")


def cell_status(value):
    """Classify a workbook cell: ('val'|'na'|'q'|'blank')."""
    if value is None:
        return "blank"
    text = str(value).strip()
    if text == "":
        return "blank"
    if text == "?":
        return "q"
    if text.lower() in NA_MARKERS:
        return "na"
    return "val"


def pair(value):
    """[display value, status] pairs keep the cell text and its interpretation side by side."""
    status = cell_status(value)
    if status == "blank":
        return [None, "blank"]
    return [value if not isinstance(value, str) else value, status]


def first_year(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if 1000 <= value <= 2100 else None
    if isinstance(value, float) and value.is_integer():
        return first_year(int(value))
    if isinstance(value, dt.date):
        return value.year
    if isinstance(value, str):
        m = re.search(r"\b(1[0-9]{3}|20[0-9]{2})\b", value)
        return int(m.group(1)) if m else None
    return None


def number(value):
    """Parse a numeric cell. Returns (number|None, was_text)."""
    if isinstance(value, bool):
        return None, False
    if isinstance(value, (int, float)):
        return float(value), False
    if isinstance(value, str):
        text = value.replace("\u00a0", " ").strip().replace(",", ".")
        try:
            return float(text), True
        except ValueError:
            return None, True
    return None, False


def facets_for(parts_text) -> list[str]:
    if cell_status(parts_text) != "val":
        return ["unclassified"]
    text = NEGATION.sub(" ", str(parts_text).lower())
    tags = [key for key, pattern in FACET_RULES.items() if re.search(pattern, text)]
    return tags or ["unclassified"]


def clean_name(value):
    return re.sub(r"\s+", " ", str(value).replace("_", " ")).strip() if value is not None else None


def load_sheet(path: Path, sheet: str | None):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        sys.exit(f"{path}: worksheet {ws.title} is empty")
    raw_headers = list(rows[0])
    headers = []
    for h in raw_headers:
        h = str(h).strip() if h is not None else None
        headers.append(HEADER_ALIASES.get(h, h) if h else None)
    labelled = [h for h in headers if h]
    width = len(headers)
    unlabelled_columns = [i for i, h in enumerate(headers) if not h]
    records = []
    for index, row in enumerate(rows[1:], start=2):
        row = list(row) + [None] * (width - len(row))
        if not any(c not in (None, "") for c in row):
            continue
        values = OrderedDict((h, row[i]) for i, h in enumerate(headers) if h)
        extra = {openpyxl.utils.get_column_letter(i + 1): row[i] for i in unlabelled_columns if row[i] not in (None, "")}
        records.append((index, values, extra))
    taxon_list = []
    if "Specimen_type_list" in wb.sheetnames:
        for r in wb["Specimen_type_list"].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                taxon_list.append({"taxon": str(r[0]).strip(), "binomial": clean_name(r[1]) if len(r) > 1 and r[1] else None})
    return ws.title, labelled, records, taxon_list


def load_previous(path: Path | None):
    if not path or not path.exists():
        return None
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".json":
        return json.loads(text)
    m = re.search(r'<script id="payload" type="application/json">(.*?)</script>', text, re.S)
    return json.loads(m.group(1)) if m else None


def build_records(records, headers, source_file, sheet_name, curated_notes):
    id_counts = Counter(v.get("Specimen_id") for _, v, _ in records)
    duplicates = {str(k): [i for i, v, _ in records if v.get("Specimen_id") == k] for k, c in id_counts.items() if c > 1}
    out = []
    for row_index, v, extra in records:
        sid = v.get("Specimen_id")
        sid_text = str(sid).strip() if sid is not None else f"row{row_index}"
        duplicate = id_counts[sid] > 1
        rec_id = f"{sid_text}@row{row_index}" if duplicate else sid_text
        notes = []
        if duplicate:
            notes.append(f"Specimen_id {sid_text} occurs in worksheet rows {', '.join(map(str, duplicates[sid_text]))}. Both records are retained; the worksheet row distinguishes them without assigning a new specimen ID.")
        if sid is None:
            notes.append("Specimen_id is blank; the worksheet row number stands in as a temporary identifier.")

        lat, lat_text = number(v.get("Latitude"))
        lon, lon_text = number(v.get("Longitude"))
        if lat_text or lon_text:
            if lat is None or lon is None:
                notes.append("A coordinate cell holds text that could not be read as a number; the record is not mapped.")
            else:
                notes.append("A coordinate cell is stored as text (with stray spacing) rather than a number. It was read numerically; check the cell format in the workbook.")
        if lat is None or lon is None:
            if cell_status(v.get("Latitude")) == "blank" and cell_status(v.get("Longitude")) == "blank":
                notes.append("No coordinates are recorded for this find; the record appears in the list and table but not on the map.")
        elif abs(lat) > 90 or abs(lon) > 180:
            notes.append(f"Coordinate out of range ({lat}, {lon}); the record is not mapped.")
            lat = lon = None

        provenance_raw = v.get("Location Provenance")
        provenance = str(provenance_raw).strip() if cell_status(provenance_raw) == "val" else None
        if provenance and slug(provenance) == "unlocated" and lat is not None:
            notes.append("Location provenance is 'unlocated' but the workbook supplies coordinates. The point is drawn as a placeholder position only.")
        if not provenance:
            notes.append("No location provenance label is recorded; the coordinate is shown without a documentation class.")

        age_raw = v.get("Age_BP_radiocarbon")
        age_text = str(age_raw) if age_raw is not None else ""
        if re.search(r"\d[,\s]+\d{4,}(?:\D|$)", age_text):
            notes.append(f"The age text \u201c{age_text.strip()}\u201d has ambiguous digit grouping. It is shown verbatim and excluded from numeric age bands until clarified.")
        if re.search(r"\b\d{1,3}\.\d{3}\b", age_text):
            notes.append(f"The age text \u201c{age_text.strip()}\u201d uses a period between digit groups. The map reads it as a thousands separator (31.700 = 31,700). Confirm the notation in the source.")

        year_found = first_year(v.get("Date_Found"))
        year_collected = first_year(v.get("Date(s)_collected"))
        if year_found is None and year_collected is None:
            notes.append("Neither a discovery year nor a collection year could be read; the record is excluded from year filters.")
        if year_found and year_collected and year_collected < year_found:
            notes.append("The collection year precedes the discovery year in the workbook.")

        species = str(v.get("Species_type") or "").strip() or "Unassigned"
        grade = GRADES.get(str(v.get("Complete?") or "").strip(), 0)
        facets = facets_for(v.get("Parts_found"))
        if grade >= 3:  # "Whole / near-whole" follows the workbook's Complete? field, not the parts wording.
            facets = [f for f in facets if f != "unclassified"] + ["carcass"]
        notes.extend(curated_notes.get(sid_text, []))
        extra_note = "; ".join(str(x).strip() for x in extra.values()) if extra else None

        rec = OrderedDict(
            id=rec_id,
            source_id=sid if not isinstance(sid, float) or not sid.is_integer() else int(sid),
            source_row=row_index,
            source_sheet=sheet_name,
            duplicate_source_id=duplicate,
            species=species,
            binomial=clean_name(v.get("Binomial")) if cell_status(v.get("Binomial")) == "val" else None,
            secondary=pair(v.get("Secondary_name(s)")),
            location=str(v.get("Location_found")).strip() if cell_status(v.get("Location_found")) == "val" else "",
            lat=lat,
            lon=lon,
            src="workbook",
            coordinate_label=provenance,
            coordinate_key=slug(provenance) if provenance else "unlabelled",
            coordinate_label_source="Location Provenance column" if provenance else None,
            geonote=None,
            lat_raw=v.get("Latitude"),
            lon_raw=v.get("Longitude"),
            date_coll=pair(v.get("Date(s)_collected")),
            date_found=pair(v.get("Date_Found")),
            age=pair(age_raw),
            complete=pair(v.get("Complete?")),
            parts=pair(v.get("Parts_found")),
            published_ref=pair(v.get("Published_References")),
            archival_ref=pair(v.get("Archival_References")),
            notes=pair(extra_note),
            common_name=pair(v.get("Common_name")),
            former_classification=pair(v.get("Former_classification")),
            year_found=year_found,
            year_collected=year_collected,
            timeline_year=year_found or year_collected,
            facets=facets,
            grade=grade,
            review_notes=notes,
            source_values=OrderedDict((k, v[k]) for k in headers) | ({"Unlabelled columns": extra} if extra else {}),
            source_file=source_file,
        )
        out.append(rec)
    return out, duplicates


def describe_changes(records, previous):
    """Compare with the previous edition's payload (ids and coordinates only)."""
    if not previous:
        return {"ids_new_to_atlas": [], "previous_ids_absent": [], "coordinate_changes": [], "previous_source_file": None}
    prev = {}
    for r in previous.get("recs", []):
        prev.setdefault(str(r.get("source_id")), []).append(r)
    current_ids = {str(r["source_id"]) for r in records}
    new_ids = sorted((i for i in current_ids if i not in prev), key=lambda s: (len(s), s))
    absent = sorted((i for i in prev if i not in current_ids), key=lambda s: (len(s), s))
    changes = []
    for r in records:
        olds = prev.get(str(r["source_id"]))
        if not olds:
            continue
        old = olds[0] if len(olds) == 1 else next((o for o in olds if o.get("common_name", [None])[0] == r["common_name"][0]), olds[0])
        if old.get("lat") is None or r["lat"] is None:
            if (old.get("lat") is None) != (r["lat"] is None):
                changes.append({"id": r["source_id"], "previous": [old.get("lat"), old.get("lon")], "workbook": [r["lat"], r["lon"]]})
            continue
        if abs(old["lat"] - r["lat"]) > 1e-4 or abs(old["lon"] - r["lon"]) > 1e-4:
            changes.append({"id": r["source_id"], "previous": [old["lat"], old["lon"]], "workbook": [r["lat"], r["lon"]]})
    return {"ids_new_to_atlas": new_ids, "previous_ids_absent": absent, "coordinate_changes": changes, "previous_source_file": previous.get("metadata", {}).get("source_file")}


def build_payload(workbook: Path, sheet: str | None, previous_path: Path | None):
    sheet_name, headers, rows, taxon_list = load_sheet(workbook, sheet)
    curated = json.loads(REVIEW_NOTES.read_text(encoding="utf-8")) if REVIEW_NOTES.exists() else {}
    curated_notes = {str(k): (v if isinstance(v, list) else [v]) for k, v in curated.get("notes", {}).items()}
    records, duplicates = build_records(rows, headers, workbook.name, sheet_name, curated_notes)
    previous = load_previous(previous_path)

    # Taxa in type-sheet order, then any others in order of first appearance.
    order = [t["taxon"] for t in taxon_list]
    present = Counter(r["species"] for r in records)
    species = OrderedDict()
    for name in order + [r["species"] for r in records]:
        if name in present and name not in species:
            species[name] = {"count": present[name]}
    years = [r["timeline_year"] for r in records if r["timeline_year"]]
    provenance = Counter(r["coordinate_label"] or "unlabelled" for r in records)
    prov_order = PROVENANCE_ORDER + sorted(k for k in provenance if k not in PROVENANCE_ORDER)
    m = re.search(r"(20\d{2})[_-](\d{2})[_-](\d{2})", workbook.name)
    source_date = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else dt.date.fromtimestamp(workbook.stat().st_mtime).isoformat()
    review = [{"source_id": r["source_id"], "row": r["source_row"], "notes": r["review_notes"]} for r in records if r["review_notes"]]
    metadata = OrderedDict(
        source_file=workbook.name,
        source_date=source_date,
        source_sheet=sheet_name,
        source_sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        source_headers=headers,
        source_rows=len(records),
        unique_specimen_ids=len({str(r["source_id"]) for r in records}),
        duplicate_ids=duplicates,
        mapped_records=sum(1 for r in records if r["lat"] is not None),
        taxon_list=taxon_list,
        coordinate_labels="Taken from the workbook's Location Provenance column.",
        provenance_counts=OrderedDict((k, provenance[k]) for k in prov_order if k in provenance),
        review_notes=review,
        built=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d"),
        **describe_changes(records, previous),
    )
    payload = OrderedDict(
        metadata=metadata,
        species=species,
        recs=records,
        W=1440,
        H=486,
        timeline={"min": min(years), "max": max(years)},
        stats={"total": len(records), "species": dict(present), "mapped": metadata["mapped_records"]},
    )
    return payload


def json_for_html(obj) -> str:
    # "</script" must never appear inside the inline JSON; escape "<" and U+2028/9 for safety.
    text = json.dumps(obj, ensure_ascii=False, separators=(",", ":"), default=str)
    return text.replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")


def render(payload, dataset_url: str, dataset_doi: str) -> str:
    md = payload["metadata"]
    total = payload["stats"]["total"]
    span = f"{payload['timeline']['min']}\u2013{payload['timeline']['max']}"
    review_items = []
    for item in md["review_notes"]:
        for note in item["notes"]:
            review_items.append(f"<li><b>#{html.escape(str(item['source_id']))}</b> (row {item['row']}): {html.escape(note)}</li>")
    changes = []
    if md.get("previous_source_file"):
        changes.append(f"<p>Compared with the previous edition (<b>{html.escape(md['previous_source_file'])}</b>): ")
        bits = []
        bits.append(f"{len(md['ids_new_to_atlas'])} ID{'s' if len(md['ids_new_to_atlas']) != 1 else ''} new to the map" + (f" ({', '.join(md['ids_new_to_atlas'])})" if md["ids_new_to_atlas"] else ""))
        bits.append(f"{len(md['previous_ids_absent'])} earlier ID{'s' if len(md['previous_ids_absent']) != 1 else ''} absent" + (f" ({', '.join(md['previous_ids_absent'])})" if md["previous_ids_absent"] else ""))
        bits.append(f"{len(md['coordinate_changes'])} record{'s' if len(md['coordinate_changes']) != 1 else ''} with changed coordinates" + (f" ({', '.join(str(c['id']) for c in md['coordinate_changes'])})" if md["coordinate_changes"] else ""))
        changes.append("; ".join(bits) + ". No equivalence, renumbering or merge has been inferred between editions; the workbook values are used as supplied.</p>")
    prov = "".join(f"<li><b>{html.escape(k)}</b>: {v} record{'s' if v != 1 else ''}</li>" for k, v in md["provenance_counts"].items())
    fills = {
        "TOTAL": str(total),
        "MAPPED": str(payload["stats"]["mapped"]),
        "UNMAPPED": str(total - payload["stats"]["mapped"]),
        "TAXA_COUNT": str(len(payload["species"])),
        "YEAR_SPAN": span,
        "YEAR_MIN": str(payload["timeline"]["min"]),
        "YEAR_MAX": str(payload["timeline"]["max"]),
        "SOURCE_FILE": html.escape(md["source_file"]),
        "SOURCE_SHEET": html.escape(md["source_sheet"]),
        "SOURCE_DATE": md["source_date"],
        "BUILD_DATE": md["built"],
        "EDITION": f"discovery-{md['source_date']}",
        "UNIQUE_IDS": str(md["unique_specimen_ids"]),
        "REVIEW_LIST": "\n".join(review_items) or "<li>No source values are flagged in this edition.</li>",
        "CHANGES": "\n".join(changes),
        "PROVENANCE_LIST": prov,
        # Inserted inside single-quoted JS strings in atlas.js.
        "DATASET_URL": dataset_url.replace("\\", "\\\\").replace("'", "\\'").replace("<", "\\u003c"),
        "DATASET_DOI": dataset_doi.replace("\\", "\\\\").replace("'", "\\'").replace("<", "\\u003c"),
    }
    def substitute(text: str, values: dict) -> str:
        missing = set(re.findall(r"{{([A-Z_]+)}}", text)) - set(values)
        if missing:
            sys.exit(f"Placeholders without values: {sorted(missing)}")
        for key, value in values.items():
            text = text.replace("{{" + key + "}}", value)
        return text

    # Small text fills first (the app script carries the dataset URL/DOI placeholders), then the
    # large embedded blocks, which are never scanned for placeholders.
    page = substitute(TEMPLATE.read_text(encoding="utf-8"), fills | {k: "{{" + k + "}}" for k in ("CSS", "APP", "LAND", "THREE", "PAYLOAD")})
    blocks = {
        "CSS": CSS.read_text(encoding="utf-8").strip(),
        "APP": substitute(APP.read_text(encoding="utf-8").strip(), fills),
        "LAND": LAND.read_text(encoding="utf-8").strip(),
        "THREE": THREE.read_text(encoding="utf-8").strip(),
        "PAYLOAD": json_for_html(payload),
    }
    for key, value in blocks.items():
        page = page.replace("{{" + key + "}}", value, 1)
    return page


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    ap.add_argument("--sheet", default=None, help="worksheet name (default: first sheet)")
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--previous", type=Path, default=None, help="previous edition (built HTML or payload JSON) for change notes")
    ap.add_argument("--payload-json", type=Path, default=None, help="also write the payload as JSON (for tests)")
    ap.add_argument("--dataset-url", default="", help="Borealis dataset landing page once published")
    ap.add_argument("--dataset-doi", default="", help="Dataset DOI once published")
    ap.add_argument("--check", action="store_true", help="print a summary and review notes without writing")
    args = ap.parse_args(argv)
    if not args.workbook.exists():
        sys.exit(f"Workbook not found: {args.workbook}")
    payload = build_payload(args.workbook, args.sheet, args.previous)
    md = payload["metadata"]
    print(f"{md['source_file']} ({md['source_sheet']}): {md['source_rows']} records, {md['unique_specimen_ids']} unique IDs, {md['mapped_records']} mapped, years {payload['timeline']['min']}-{payload['timeline']['max']}")
    print("Taxa:", ", ".join(f"{k} {v['count']}" for k, v in payload["species"].items()))
    print("Provenance:", ", ".join(f"{k} {v}" for k, v in md["provenance_counts"].items()))
    if md.get("previous_source_file"):
        print(f"vs {md['previous_source_file']}: new {md['ids_new_to_atlas']}, absent {md['previous_ids_absent']}, coordinate changes {[c['id'] for c in md['coordinate_changes']]}")
    for item in md["review_notes"]:
        for note in item["notes"]:
            print(f"  review #{item['source_id']} (row {item['row']}): {note}")
    if args.check:
        return
    if args.payload_json:
        args.payload_json.parent.mkdir(parents=True, exist_ok=True)
        args.payload_json.write_text(json.dumps(payload, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    args.out.write_text(render(payload, args.dataset_url, args.dataset_doi), encoding="utf-8")
    print(f"Wrote {args.out} ({args.out.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
